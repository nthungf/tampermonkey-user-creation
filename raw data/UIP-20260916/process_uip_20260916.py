import os
import csv
import json
import openpyxl
from openpyxl import Workbook

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(CURRENT_DIR, "UIP List-20260910.xlsx")

wb = openpyxl.load_workbook(EXCEL_PATH)
sheet = wb['Users']

headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]

existing_dont_do = []
existing_need_deleted = []
not_exist_need_to_add = []
all_emails = []

for r in range(2, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
    if not any(vals):
        continue
    row_dict = dict(zip(headers, vals))
    
    # Strip string fields
    clean_dict = {}
    for k, v in row_dict.items():
        if isinstance(v, str):
            clean_dict[k] = v.strip()
        else:
            clean_dict[k] = v
            
    email = str(clean_dict.get('Email') or '').strip()
    if email:
        all_emails.append(email)
        
    note = clean_dict.get('Notes')
    if note == 'Delete account':
        existing_need_deleted.append(clean_dict)
    elif note in ['Add thêm NV mới', 'Add thêm']:
        not_exist_need_to_add.append(clean_dict)
    elif note is None or str(note).strip() == '':
        existing_dont_do.append(clean_dict)
    else:
        print(f"Warning: Unknown note: {note} for row {vals}")

print(f"Processed summary:")
print(f"  Total emails: {len(all_emails)}")
print(f"  1. existing_and_dont_do: {len(existing_dont_do)}")
print(f"  2. existing_and_need_deleted: {len(existing_need_deleted)}")
print(f"  3. not_exist_need_to_add: {len(not_exist_need_to_add)}")

# Function to save CSV (UTF-8-SIG for Excel compatibility)
def save_csv(filename, data):
    filepath = os.path.join(CURRENT_DIR, filename)
    with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in data:
            writer.writerow(row)
    print(f"Saved {filepath}")

# Function to save XLSX
def save_xlsx(filename, data):
    filepath = os.path.join(CURRENT_DIR, filename)
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Users"
    out_ws.append(headers)
    for row in data:
        out_ws.append([row.get(h) for h in headers])
    out_wb.save(filepath)
    print(f"Saved {filepath}")

# 1. Save CSV files
save_csv("existing_and_dont_do.csv", existing_dont_do)
save_csv("existing_and_need_deleted.csv", existing_need_deleted)
save_csv("not_exist_need_to_add.csv", not_exist_need_to_add)

# 2. Save XLSX files
save_xlsx("existing_and_dont_do.xlsx", existing_dont_do)
save_xlsx("existing_and_need_deleted.xlsx", existing_need_deleted)
save_xlsx("not_exist_need_to_add.xlsx", not_exist_need_to_add)

# 3. Generate SQL query file
sql_filepath = os.path.join(CURRENT_DIR, "check_emails_exist.sql")
with open(sql_filepath, 'w', encoding='utf-8') as f:
    f.write("-- =============================================================================\n")
    f.write("-- SQL Queries to check existing & non-existing emails for United Pharma (UIP)\n")
    f.write(f"-- Source: UIP List-20260910.xlsx (Total {len(all_emails)} emails)\n")
    f.write("-- Company ID: 1237\n")
    f.write("-- =============================================================================\n\n")

    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("-- QUERY 1: CTE matching all input emails against DB to see status of each email\n")
    f.write("-- (Status: 'EXIST' vs 'NOT EXIST')\n")
    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("WITH input_emails AS (\n")
    for i, email in enumerate(all_emails):
        prefix = "    SELECT " if i == 0 else "    UNION ALL SELECT "
        f.write(f"{prefix}'{email}' AS email\n")
    f.write(")\n")
    f.write("SELECT \n")
    f.write("    i.email AS input_email,\n")
    f.write("    CASE \n")
    f.write("        WHEN u.id IS NOT NULL THEN 'EXIST'\n")
    f.write("        ELSE 'NOT EXIST'\n")
    f.write("    END AS check_status,\n")
    f.write("    u.id AS user_id,\n")
    f.write("    u.employee_code,\n")
    f.write("    u.name,\n")
    f.write("    department.name AS department_name,\n")
    f.write("    job_title.name AS job_title_name,\n")
    f.write("    u.status AS active_status\n")
    f.write("FROM input_emails i\n")
    f.write("LEFT JOIN user u ON LOWER(TRIM(i.email)) = LOWER(TRIM(u.email)) AND u.company_id = 1237\n")
    f.write("LEFT JOIN department ON u.department_id = department.id\n")
    f.write("LEFT JOIN job_title ON u.job_title_id = job_title.id\n")
    f.write("ORDER BY check_status ASC, i.email ASC;\n\n")

    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("-- QUERY 2: Filter ONLY emails that do NOT exist in DB (Need to be created new)\n")
    f.write("-- Expected output: 4 emails to create\n")
    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("WITH input_emails AS (\n")
    for i, email in enumerate(all_emails):
        prefix = "    SELECT " if i == 0 else "    UNION ALL SELECT "
        f.write(f"{prefix}'{email}' AS email\n")
    f.write(")\n")
    f.write("SELECT \n")
    f.write("    i.email AS missing_email\n")
    f.write("FROM input_emails i\n")
    f.write("LEFT JOIN user u ON LOWER(TRIM(i.email)) = LOWER(TRIM(u.email)) AND u.company_id = 1237\n")
    f.write("WHERE u.id IS NULL;\n\n")

    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("-- QUERY 3: Simple SELECT for all emails in list that EXIST in database\n")
    f.write("-- -----------------------------------------------------------------------------\n")
    email_list_str = ",\n    ".join(f"'{e}'" for e in all_emails)
    f.write("SELECT \n")
    f.write("    u.id,\n")
    f.write("    u.employee_code,\n")
    f.write("    u.name,\n")
    f.write("    u.email,\n")
    f.write("    department.name AS department_name,\n")
    f.write("    job_title.name AS job_title_name,\n")
    f.write("    u.status\n")
    f.write("FROM user u\n")
    f.write("LEFT JOIN department ON u.department_id = department.id\n")
    f.write("LEFT JOIN job_title ON u.job_title_id = job_title.id\n")
    f.write("WHERE u.company_id = 1237\n")
    f.write(f"  AND u.email IN (\n    {email_list_str}\n  );\n\n")

    f.write("-- -----------------------------------------------------------------------------\n")
    f.write("-- QUERY 4: Deactivate/Delete the 16 users marked 'Delete account'\n")
    f.write("-- -----------------------------------------------------------------------------\n")
    del_emails = [str(r.get('Email')).strip() for r in existing_need_deleted]
    del_emails_str = ",\n    ".join(f"'{e}'" for e in del_emails)
    f.write("-- Check details of 16 accounts before deactivation:\n")
    f.write("SELECT u.id, u.employee_code, u.name, u.email, u.status\n")
    f.write("FROM user u\n")
    f.write("WHERE u.company_id = 1237\n")
    f.write(f"  AND u.email IN (\n    {del_emails_str}\n  );\n\n")
    f.write("-- Deactivate (soft delete status = 0):\n")
    f.write("UPDATE user\n")
    f.write("SET status = 0\n")
    f.write("WHERE company_id = 1237\n")
    f.write(f"  AND email IN (\n    {del_emails_str}\n  );\n")

print(f"Saved {sql_filepath}")

# 4. Generate users_ready_for_bot.json for the 4 users to be added
bot_users = []
for u in not_exist_need_to_add:
    emp_code = str(u.get('Mã nhân viên') or '').strip()
    clean_phone = str(u.get('SĐT') or '').strip().replace(" ", "")
    bot_users.append({
        "name": str(u.get('Họ tên') or '').strip(),
        "email": str(u.get('Email') or '').strip(),
        "password": clean_phone if clean_phone else "123456",
        "phone": clean_phone,
        "employee_code": emp_code,
        "gender": str(u.get('Giới tính') or '').strip(),
        "limit": "",
        "department": str(u.get('Phòng ban') or '').strip(),
        "branch": "",
        "job_title": str(u.get('Chức vụ') or '').strip(),
        "role": str(u.get('Vai trò') or 'Thành viên').strip(),
        "note": str(u.get('Notes') or '').strip(),
        "level": "",
        "approving_manager": ""
    })

bot_json_path = os.path.join(CURRENT_DIR, "users_ready_for_bot.json")
with open(bot_json_path, 'w', encoding='utf-8') as f:
    json.dump(bot_users, f, ensure_ascii=False, indent=2)
print(f"Saved {bot_json_path}")
