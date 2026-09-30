import openpyxl
import pandas as pd
import json
import os

def main():
    excel_path = '/home/hung/Desktop/explore/tampermonkey-user-creation/raw data/dong-a-20260904/Phân quyền Mytour Booking- Ethical Sep\'2026.xlsx'
    out_dir = '/home/hung/Desktop/explore/tampermonkey-user-creation/raw data/dong-a-20260904'
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_name = 'Update Mytour 03.09.2026'
    sheet = wb[sheet_name]

    # 1. Extract CSV
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(list(row))

    df = pd.DataFrame(data).dropna(how='all').dropna(axis=1, how='all')

    csv_clean = os.path.join(out_dir, f'{sheet_name}.csv')
    csv_full = os.path.join(out_dir, f"Phân quyền Mytour Booking- Ethical Sep'2026_{sheet_name}.csv")

    df.to_csv(csv_clean, index=False, header=False, encoding='utf-8-sig')
    df.to_csv(csv_full, index=False, header=False, encoding='utf-8-sig')
    print(f"Extracted CSV to: {csv_clean}")

    # 2. Parse 4 tables
    tables = []
    current_header = None
    current_rows = []

    for r in range(1, sheet.max_row + 1):
        val0 = sheet.cell(r, 1).value
        val1 = sheet.cell(r, 2).value
        row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        
        if val0 and any(h in str(val0).upper() for h in ['ETHICAL-']):
            if current_rows:
                tables.append({'header': current_header, 'rows': current_rows})
                current_rows = []
            current_header = str(val0).strip()
        elif any(v is not None for v in row_vals):
            if val0 == 'Ref' or (val1 == 'StaffID'):
                continue
            if val0 is not None and str(val0).isdigit():
                current_rows.append(row_vals)

    if current_rows:
        tables.append({'header': current_header, 'rows': current_rows})

    print(f"Loaded {len(tables)} tables.")

    # 3. Process newcomers (Ref != 1)
    newcomers = []
    for t in tables:
        rows = t['rows']
        if not rows:
            continue
        approver_row = rows[0]
        approver_name = str(approver_row[2]).strip()
        approver_email = str(approver_row[11]).strip()

        for r in rows[1:]: # newcomer rows
            ref = r[0]
            staff_id = str(r[1]).strip()
            name = str(r[2]).strip()
            lc_with = str(r[3]).strip().upper() if r[3] else ""
            dept = str(r[4]).strip() if r[4] else ""
            group = str(r[5]).strip() if r[5] else ""
            job = str(r[6]).strip() if r[6] else ""
            role_desc = str(r[7]).strip() if r[7] else ""
            class_a = r[8]
            class_b = r[9]
            class_c = r[10]
            email = str(r[11]).strip()

            # Normalize typo in Job Title: 'Key Account Excutive' -> 'Key Account Executive'
            if job.lower() == 'key account excutive':
                job = 'Key Account Executive'

            # Department assignment based on LC With
            if 'MAXXCARE' in lc_with:
                department = 'Ethical | Maxxcare'
            else:
                department = 'Ethical | HD'

            newcomers.append({
                "name": name,
                "email": email,
                "password": f"Mwc@{staff_id}",
                "phone": "0123456789",
                "employee_code": staff_id,
                "gender": "",
                "limit": "",
                "department": department,
                "branch": "",
                "job_title": job,
                "role": "Thành viên",
                "note": "",
                "level": "",
                "approving_manager": approver_name,
                "approving_manager_email": approver_email,
                "class_c": class_c,
                "table": t['header']
            })

    print(f"Total newcomers: {len(newcomers)}")

    # 4. Generate users_ready_for_bot.json
    users_bot = []
    for u in newcomers:
        users_bot.append({
            "name": u["name"],
            "email": u["email"],
            "password": u["password"],
            "phone": u["phone"],
            "employee_code": u["employee_code"],
            "gender": u["gender"],
            "limit": u["limit"],
            "department": u["department"],
            "branch": u["branch"],
            "job_title": u["job_title"],
            "role": u["role"],
            "note": u["note"],
            "level": u["level"],
            "approving_manager": u["approving_manager"]
        })

    bot_json_path = os.path.join(out_dir, "users_ready_for_bot.json")
    with open(bot_json_path, 'w', encoding='utf-8') as f:
        json.dump(users_bot, f, indent=2, ensure_ascii=False)
    print(f"Saved users_ready_for_bot.json to: {bot_json_path}")

    # 5. Generate individual_approval_flows.json
    approval_flows = []
    for u in newcomers:
        flow = {
            "target_user_emails": [u["email"]],
            "approval_steps": [
                {
                    "step_number": 1,
                    "approver_email": u["approving_manager_email"]
                },
                {
                    "step_number": 2,
                    "approver_email": "adminoffice@megawecare.com"
                }
            ]
        }
        approval_flows.append(flow)

    flow_json_path = os.path.join(out_dir, "individual_approval_flows.json")
    with open(flow_json_path, 'w', encoding='utf-8') as f:
        json.dump(approval_flows, f, indent=2, ensure_ascii=False)
    print(f"Saved individual_approval_flows.json to: {flow_json_path}")

    # 6. Generate target_emails.json
    budget_map = {}
    for u in newcomers:
        c_val = u["class_c"]
        if pd.notna(c_val) and c_val != "":
            try:
                val_str = str(int(float(c_val) // 1000))
            except:
                val_str = str(c_val)
        else:
            val_str = "Unknown"

        if val_str not in budget_map:
            budget_map[val_str] = []
        budget_map[val_str].append(u["email"])

    sorted_budget_map = dict(sorted(budget_map.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 0, reverse=True))

    budget_json_path = os.path.join(out_dir, "target_emails.json")
    with open(budget_json_path, 'w', encoding='utf-8') as f:
        json.dump(sorted_budget_map, f, indent=2, ensure_ascii=False)
    print(f"Saved target_emails.json to: {budget_json_path}")

if __name__ == '__main__':
    main()
